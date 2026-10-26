import io

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from accounts.models import Student
from portal import masterlist

User = get_user_model()

CSV = (
    "Student No.,Last Name,First Name,Middle Name,Year Level,Section\n"
    "2024-0001,Dela Cruz,Juan,Santos,1st Year,A\n"
    "2024-0002,Reyes,Maria,,2,B\n"
    ",NoNumber,Ana,,1,A\n"
    "2024-0001,Dup,Juan,,1,A\n"
    "2024-0003,Lim,Bob,,9,A\n"
)


def _upload(name, data):
    return SimpleUploadedFile(name, data)


class MasterlistParsingTests(TestCase):
    def test_csv_validation_flags_each_problem(self):
        rows = masterlist.validate_rows(masterlist.parse_upload(_upload("m.csv", CSV.encode())))
        status = [r["status"] for r in rows]
        self.assertEqual(status, ["new", "new", "error", "error", "error"])
        self.assertEqual(rows[0]["year_level"], "1")
        self.assertIn("Duplicate", rows[3]["message"])
        self.assertIn("not 1-4", rows[4]["message"])

    def test_xlsx_and_numeric_cells(self):
        wb = Workbook()
        ws = wb.active
        ws.append(masterlist.COLUMNS)
        ws.append([20240005, "Tan", "Eve", None, 3, "C"])
        buf = io.BytesIO()
        wb.save(buf)
        rows = masterlist.validate_rows(masterlist.parse_upload(_upload("m.xlsx", buf.getvalue())))
        self.assertEqual(rows[0]["student_number"], "20240005")
        self.assertEqual(rows[0]["year_level"], "3")
        self.assertEqual(rows[0]["status"], "new")

    def test_missing_columns_and_bad_type(self):
        with self.assertRaises(masterlist.MasterlistError):
            masterlist.parse_upload(_upload("m.csv", b"Foo,Bar\n1,2\n"))
        with self.assertRaises(masterlist.MasterlistError):
            masterlist.parse_upload(_upload("m.pdf", b"x"))

    def test_templates_round_trip(self):
        self.assertIn("Student No.", masterlist.template_csv_bytes().decode("utf-8-sig"))
        wb = load_workbook(io.BytesIO(masterlist.template_xlsx_bytes()))
        self.assertEqual([c.value for c in wb["Masterlist"][1]], masterlist.COLUMNS)

    def test_import_creates_accounts_that_can_log_in_and_skips_existing(self):
        rows = masterlist.validate_rows(masterlist.parse_upload(_upload("m.csv", CSV.encode())))
        creds = masterlist.import_rows(rows)
        self.assertEqual(len(creds), 2)
        user = User.objects.get(username="2024-0001")
        self.assertTrue(user.is_student)
        self.assertTrue(user.check_password(creds[0]["password"]))
        self.assertEqual(Student.objects.get(user=user).year_level, "1")
        # second pass: both now exist
        again = masterlist.validate_rows(masterlist.parse_upload(_upload("m.csv", CSV.encode())))
        self.assertEqual(masterlist.summarize(again)["exists"], 2)
        self.assertEqual(masterlist.import_rows(again), [])


HEADER = "Student No.,Last Name,First Name,Middle Name,Year Level,Section\n"


def _rows(text):
    return masterlist.validate_rows(masterlist.parse_upload(_upload("m.csv", (HEADER + text).encode())))


class MasterlistUpdateTests(TestCase):
    """The Student No. is the key: a registered number is updated, never duplicated."""

    def setUp(self):
        creds = masterlist.import_rows(_rows("2024-0001,Dela Cruz,Juan,Santos,1,A\n2024-0002,Reyes,Maria,,2,B\n"))
        self.password = {c["student_number"]: c["password"] for c in creds}

    def student(self, number):
        return Student.objects.select_related("user").get(student_number=number)

    def test_a_changed_section_and_year_is_an_update_that_keeps_the_password(self):
        rows = _rows("2024-0001,Dela Cruz,Juan,Santos,2,C\n2024-0002,Reyes,Maria,,2,B\n")
        self.assertEqual([r["status"] for r in rows], ["update", "exists"])
        self.assertEqual(sorted(rows[0]["changes"]), ["Section: A -> C", "Year: 1 -> 2"])
        self.assertEqual(masterlist.import_rows(rows), [])                     # nothing created, no new slips
        s = self.student("2024-0001")
        self.assertEqual((s.year_level, s.section), ("2", "C"))
        self.assertTrue(s.user.check_password(self.password["2024-0001"]))     # same password still works
        self.assertEqual(Student.objects.count(), 2)                           # never a second record

    def test_names_can_be_corrected(self):
        rows = _rows("2024-0002,Reyes-Santos,Maria Clara,Lopez,2,B\n")
        self.assertEqual(rows[0]["status"], "update")
        masterlist.import_rows(rows)
        u = self.student("2024-0002").user
        self.assertEqual((u.last_name, u.first_name, u.middle_name), ("Reyes-Santos", "Maria Clara", "Lopez"))

    def test_blank_cells_never_erase_stored_values(self):
        rows = _rows("2024-0001,Dela Cruz,Juan,,,\n")
        self.assertEqual(rows[0]["status"], "exists")                          # nothing differs
        masterlist.import_rows(rows)
        s = self.student("2024-0001")
        self.assertEqual((s.user.middle_name, s.year_level, s.section), ("Santos", "1", "A"))

    def test_the_number_matches_regardless_of_case(self):
        masterlist.import_rows(_rows("ABC-1,Lim,Bob,,1,A\n"))
        rows = _rows("abc-1,Lim,Bob,,3,D\n")
        self.assertEqual(rows[0]["status"], "update")
        masterlist.import_rows(rows)
        self.assertEqual(Student.objects.filter(student_number__iexact="abc-1").count(), 1)
        self.assertEqual(self.student("ABC-1").section, "D")

    def test_a_row_with_an_invalid_value_updates_nothing(self):
        rows = _rows("2024-0001,Dela Cruz,Juan,Santos,9,Z\n")
        self.assertEqual(rows[0]["status"], "error")
        masterlist.import_rows(rows)
        self.assertEqual((self.student("2024-0001").year_level, self.student("2024-0001").section), ("1", "A"))

    def test_a_number_twice_in_one_file_is_recorded_once(self):
        rows = _rows("2024-0050,Tan,Eve,,1,A\n2024-0050,Tan,Eve,,2,B\n")
        self.assertEqual([r["status"] for r in rows], ["new", "error"])
        self.assertIn("Duplicate", rows[1]["message"])
        self.assertEqual(len(masterlist.import_rows(rows)), 1)
        s = self.student("2024-0050")
        self.assertEqual((s.year_level, s.section), ("1", "A"))                # the first row won

    def test_updating_the_same_number_twice_in_a_file_does_not_apply_the_second(self):
        rows = _rows("2024-0001,Dela Cruz,Juan,Santos,3,B\n2024-0001,Dela Cruz,Juan,Santos,4,C\n")
        self.assertEqual([r["status"] for r in rows], ["update", "error"])
        masterlist.import_rows(rows)
        self.assertEqual((self.student("2024-0001").year_level, self.student("2024-0001").section), ("3", "B"))

    def test_portal_preview_shows_the_changes_and_confirm_applies_them(self):
        User.objects.create_user("adviser", password="pw-12345-x", is_admin=True)
        self.client.login(username="adviser", password="pw-12345-x")
        csv = (HEADER + "2024-0001,Dela Cruz,Juan,Santos,2,C\n2024-0009,New,Student,,1,A\n").encode()
        r = self.client.post(reverse("portal:students-import"), {"action": "preview", "file": _upload("m.csv", csv)})
        self.assertContains(r, "1 to update")
        self.assertContains(r, "Section: A -&gt; C")
        self.assertEqual(self.student("2024-0001").section, "A")               # preview changes nothing
        r = self.client.post(reverse("portal:students-import"), {"action": "confirm"}, follow=True)
        self.assertEqual(self.student("2024-0001").section, "C")
        self.assertTrue(Student.objects.filter(student_number="2024-0009").exists())
        self.assertContains(r, "Created 1 student(s), updated 1")


class MasterlistErrorReportTests(TestCase):
    def setUp(self):
        User.objects.create_user("adviser", password="pw-12345-x", is_admin=True)
        self.client.login(username="adviser", password="pw-12345-x")
        self.csv = (HEADER + "2025-1,Good,One,,1,A\n,NoNumber,Two,,1,A\n2025-3,Good,Three,,2,B\n"
                    "2025-4,Bad,Year,,9,A\n" + "X" * 60 + ",Long,Number,,1,A\n").encode()

    def preview(self):
        return self.client.post(reverse("portal:students-import"),
                                {"action": "preview", "file": _upload("m.csv", self.csv)})

    def test_review_order_puts_errors_first_then_file_order(self):
        rows = _rows("2025-1,Good,One,,1,A\n,NoNumber,Two,,1,A\n2025-3,Good,Three,,2,B\n2025-4,Bad,Year,,9,A\n")
        self.assertEqual([r["line"] for r in masterlist.review_order(rows)], [3, 5, 2, 4])

    def test_too_long_student_number_has_its_own_message(self):
        self.assertEqual(_rows("X" * 60 + ",Long,Number,,1,A\n")[0]["message"],
                         "Student No. is longer than 50 characters")

    def test_preview_lists_errors_first_with_a_warning_and_a_download_button(self):
        r = self.preview()
        self.assertContains(r, "3 rows with errors will be skipped")
        self.assertContains(r, "the other rows are still imported")
        self.assertContains(r, reverse("portal:students-import-errors"))
        html = r.content.decode()
        self.assertLess(html.index("Student No. is required"), html.index("Three"))     # errors before good rows
        self.assertEqual(html.count('class="table-danger"'), 3)

    def test_download_gives_only_the_error_rows_with_line_and_problem(self):
        self.preview()
        r = self.client.get(reverse("portal:students-import-errors"))
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r["Content-Disposition"])
        text = r.content.decode("utf-8-sig")
        lines = text.strip().splitlines()
        self.assertEqual(lines[0], "Student No.,Last Name,First Name,Middle Name,Year Level,Section,Line,Problem")
        self.assertEqual(len(lines), 4)                                                 # header + 3 errors
        self.assertIn("Year Level '9' is not 1-4", text)
        self.assertNotIn("Three", text)                                                 # good rows are not in it

    def test_the_downloaded_file_can_be_uploaded_again(self):
        self.preview()
        csv_bytes = self.client.get(reverse("portal:students-import-errors")).content
        rows = masterlist.validate_rows(masterlist.parse_upload(_upload("fix.csv", csv_bytes)))
        self.assertEqual(len(rows), 3)                                                  # extra Line / Problem columns ignored
        self.assertTrue(all(r["status"] == "error" for r in rows))

    def test_excel_download_has_the_same_rows_and_keeps_ids_as_text(self):
        csv_bytes = (HEADER + "0012345,Lead,Zero,,9,A\n2025-2,Fine,Row,,1,A\n").encode()
        self.client.post(reverse("portal:students-import"), {"action": "preview", "file": _upload("m.csv", csv_bytes)})
        r = self.client.get(reverse("portal:students-import-errors") + "?format=xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn(".xlsx", r["Content-Disposition"])
        ws = load_workbook(io.BytesIO(r.content))["Masterlist"]
        self.assertEqual([c.value for c in ws[1]], masterlist.ERROR_REPORT_COLUMNS)
        self.assertEqual(ws.max_row, 2)                                               # header + the one error row
        self.assertEqual(ws["A2"].value, "0012345")                                   # leading zeros survive
        self.assertIn("is not 1-4", ws["H2"].value)

    def test_the_excel_error_file_can_be_uploaded_again(self):
        self.preview()
        data = self.client.get(reverse("portal:students-import-errors") + "?format=xlsx").content
        rows = masterlist.validate_rows(masterlist.parse_upload(_upload("fix.xlsx", data)))
        self.assertEqual(len(rows), 3)
        self.assertTrue(all(r["status"] == "error" for r in rows))

    def test_preview_offers_excel_and_csv(self):
        r = self.preview()
        self.assertContains(r, reverse("portal:students-import-errors") + "?format=xlsx")
        self.assertContains(r, "Download them (Excel)")

    def test_no_errors_means_no_download(self):
        self.client.post(reverse("portal:students-import"),
                         {"action": "preview", "file": _upload("m.csv", (HEADER + "2025-9,Fine,Row,,1,A\n").encode())})
        r = self.client.get(reverse("portal:students-import-errors"), follow=True)
        self.assertContains(r, "no rows with errors")

    def test_download_needs_an_adviser(self):
        self.client.logout()
        self.assertEqual(self.client.get(reverse("portal:students-import-errors")).status_code, 302)


class MasterlistPortalTests(TestCase):
    def setUp(self):
        User.objects.create_user("adviser", password="pw-12345-x", is_admin=True)
        self.client.login(username="adviser", password="pw-12345-x")

    def test_pages_and_templates_load(self):
        for name in ("students-import", "students-credentials"):
            self.assertEqual(self.client.get(reverse(f"portal:{name}")).status_code, 200)
        for fmt in ("csv", "xlsx"):
            r = self.client.get(reverse("portal:students-import-template", args=[fmt]))
            self.assertEqual(r.status_code, 200)
            self.assertIn("attachment", r["Content-Disposition"])
        self.assertEqual(self.client.get(reverse("portal:students-import-template", args=["pdf"])).status_code, 404)
        self.assertContains(self.client.get(reverse("portal:students-list")), "Import masterlist")

    def test_preview_confirm_print_and_csv(self):
        r = self.client.post(reverse("portal:students-import"),
                             {"action": "preview", "file": _upload("m.csv", CSV.encode())})
        self.assertContains(r, "2 new")
        self.assertEqual(Student.objects.count(), 0, list(Student.objects.values_list("student_number", flat=True)))
        r = self.client.post(reverse("portal:students-import"), {"action": "confirm"}, follow=True)
        self.assertEqual(Student.objects.count(), 2, list(Student.objects.values_list("student_number", flat=True)))
        self.assertContains(r, "Dela Cruz")
        self.assertContains(r, "Password")
        csv_resp = self.client.get(reverse("portal:students-credentials") + "?format=csv")
        self.assertIn("2024-0002", csv_resp.content.decode())
        self.client.post(reverse("portal:students-credentials") + "?clear=1")
        self.assertNotContains(self.client.get(reverse("portal:students-credentials")), "Dela Cruz")

    def test_bulk_reset_changes_passwords_for_the_filter_only(self):
        a = User.objects.create_user("s1", password="old-pass-1", is_student=True)
        b = User.objects.create_user("s2", password="old-pass-2", is_student=True)
        Student.objects.create(user=a, student_number="s1", year_level="1", section="A")
        Student.objects.create(user=b, student_number="s2", year_level="2", section="A")
        url = reverse("portal:students-credentials-reset")
        self.client.post(url, {"year_level": "1"})                  # no confirmation -> nothing
        a.refresh_from_db()
        self.assertTrue(a.check_password("old-pass-1"))
        r = self.client.post(url, {"year_level": "1", "confirm": "yes"}, follow=True)
        a.refresh_from_db(); b.refresh_from_db()
        self.assertFalse(a.check_password("old-pass-1"))
        self.assertTrue(b.check_password("old-pass-2"))
        self.assertContains(r, "s1")

    def test_non_admin_is_refused(self):
        self.client.logout()
        User.objects.create_user("stu", password="pw-12345-x", is_student=True)
        self.client.login(username="stu", password="pw-12345-x")
        self.assertEqual(self.client.get(reverse("portal:students-import")).status_code, 403)


class LogsAndNavTests(TestCase):
    def setUp(self):
        User.objects.create_user("adviser", password="pw-12345-x", is_admin=True)
        self.client.login(username="adviser", password="pw-12345-x")

    def test_logs_csv_download(self):
        r = self.client.get(reverse("portal:attendance_logs-list") + "?format=csv")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "text/csv")
        self.assertIn("Student No.", r.content.decode())

    def test_logs_page_has_download_and_print(self):
        r = self.client.get(reverse("portal:attendance_logs-list"))
        self.assertContains(r, "Download CSV")
        self.assertContains(r, "window.print()")

    def test_nav_groups(self):
        from portal.context import sidebar_nav

        nav = {i["label"]: i for i in sidebar_nav(None)["nav_items"]}
        self.assertEqual([c["label"] for c in nav["Reports"]["children"]],
                         ["Attendance Report", "Attendance Logs", "Financial Report", "Fines"])
        self.assertEqual([c["label"] for c in nav["Settings"]["children"]],
                         ["Appearance", "School Years & Semesters"])
        self.assertNotIn("School Years", nav)


class TermsPageTests(TestCase):
    def setUp(self):
        User.objects.create_user("adviser", password="pw-12345-x", is_admin=True)
        self.client.login(username="adviser", password="pw-12345-x")

    def test_one_page_two_add_buttons_nested_rows(self):
        from attendance.models import SchoolYear, Semester

        sy = SchoolYear.objects.create(sy="2030-2031")
        Semester.objects.create(school_year=sy, name="1st")
        SchoolYear.objects.create(sy="2029-2030")
        r = self.client.get(reverse("portal:terms-list"))
        self.assertContains(r, "Add School Year")
        self.assertContains(r, "Add Semester")
        html = r.content.decode()
        # one flat row per semester: school year and semester in the same <tr>
        self.assertRegex(html, r"(?s)<tr>\s*<td>\s*2030-2031.*?</td>\s*<td>\s*1st Semester")
        self.assertIn("2029-2030", html)             # a year without semesters still has a row
        self.assertIn("No semester yet", html)
        # the old list URLs show the same combined page; the table fragment refreshes via HTMX
        self.assertContains(self.client.get(reverse("portal:semesters-list")), "Add School Year")
        part = self.client.get(reverse("portal:terms-list") + "?partial=table", HTTP_HX_REQUEST="true")
        self.assertNotContains(part, "<html")
        self.assertContains(part, "1st Semester")

    def test_add_forms_still_work(self):
        r = self.client.post(reverse("portal:school_years-add"), {"sy": "2031-2032"}, HTTP_HX_REQUEST="true")
        self.assertEqual(r.status_code, 204, r.content[:200])
        from attendance.models import SchoolYear
        sy = SchoolYear.objects.get(sy="2031-2032")
        r = self.client.post(reverse("portal:semesters-add"), {"school_year": sy.pk, "name": "2nd"}, HTTP_HX_REQUEST="true")
        self.assertEqual(r.status_code, 204, r.content[:200])
        self.assertEqual(sy.semesters.count(), 1)


class TermsLockedWhenUsedTests(TestCase):
    """Once an event exists, only the Super Admin can edit/delete its school year / semester."""

    def setUp(self):
        from attendance.models import SchoolYear, Semester
        from attendance.tests import make_event

        self.adv = User.objects.create_user("adviser", password="pw-12345-x", is_admin=True)
        self.boss = User.objects.create_user("boss", password="pw-12345-x", is_admin=True,
                                             is_superuser=True, is_staff=True)
        self.event = make_event()                       # creates 2025-2026 / 1st with one event
        self.sem = self.event.semester
        self.sy = self.sem.school_year
        self.free_sy = SchoolYear.objects.create(sy="2099-2100")
        self.free_sem = Semester.objects.create(school_year=self.free_sy, name="2nd")

    def _urls(self):
        return {
            "sem_edit": reverse("portal:semesters-edit", args=[self.sem.pk]),
            "sem_del": reverse("portal:semesters-delete", args=[self.sem.pk]),
            "sy_edit": reverse("portal:school_years-edit", args=[self.sy.pk]),
            "sy_del": reverse("portal:school_years-delete", args=[self.sy.pk]),
        }

    def test_buttons_hidden_for_adviser_but_shown_for_super_admin(self):
        u = self._urls()
        self.client.force_login(self.adv)
        html = self.client.get(reverse("portal:terms-list")).content.decode()
        for key in u.values():
            self.assertNotIn(key, html)
        self.assertIn("Locked", html)
        # a school year / semester without events stays editable
        self.assertIn(reverse("portal:semesters-edit", args=[self.free_sem.pk]), html)
        self.assertIn(reverse("portal:school_years-edit", args=[self.free_sy.pk]), html)
        self.client.force_login(self.boss)
        html = self.client.get(reverse("portal:terms-list")).content.decode()
        for key in u.values():
            self.assertIn(key, html)

    def test_adviser_is_refused_server_side(self):
        from attendance.models import SchoolYear, Semester

        u = self._urls()
        self.client.force_login(self.adv)
        for name in ("sem_edit", "sy_edit"):
            r = self.client.get(u[name], HTTP_HX_REQUEST="true")
            self.assertEqual(r.status_code, 204)
            self.assertIn("Super Admin", r["HX-Trigger"])
        self.client.post(u["sem_del"], HTTP_HX_REQUEST="true")
        self.client.post(u["sy_del"], HTTP_HX_REQUEST="true")
        self.client.post(u["sy_edit"], {"sy": "1999-2000"}, HTTP_HX_REQUEST="true")
        self.assertTrue(Semester.objects.filter(pk=self.sem.pk).exists())
        self.assertEqual(SchoolYear.objects.get(pk=self.sy.pk).sy, "2025-2026")
        # unlocked records can still be deleted by the adviser
        self.assertEqual(self.client.post(reverse("portal:semesters-delete", args=[self.free_sem.pk]),
                                          HTTP_HX_REQUEST="true").status_code, 204)
        self.assertFalse(Semester.objects.filter(pk=self.free_sem.pk).exists())

    def test_api_blocks_adviser_but_not_super_admin(self):
        from rest_framework.test import APIClient

        api = APIClient()
        api.force_authenticate(self.adv)
        self.assertEqual(api.patch(f"/api/school-years/{self.sy.pk}/", {"sy": "1999-2000"}, format="json").status_code, 403)
        self.assertEqual(api.delete(f"/api/semesters/{self.sem.pk}/").status_code, 403)
        self.assertEqual(api.delete(f"/api/semesters/{self.free_sem.pk}/").status_code, 204)
        sup = APIClient()
        sup.force_authenticate(self.boss)
        self.assertEqual(sup.patch(f"/api/school-years/{self.sy.pk}/", {"sy": "1999-2000"}, format="json").status_code, 200)

    def test_super_admin_can_still_delete_through_the_portal(self):
        from attendance.models import Semester

        self.client.force_login(self.boss)
        r = self.client.post(self._urls()["sem_del"], HTTP_HX_REQUEST="true")
        self.assertEqual(r.status_code, 204)
        self.assertFalse(Semester.objects.filter(pk=self.sem.pk).exists())
