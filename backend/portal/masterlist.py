"""Student masterlist import (CSV / Excel) and bulk credential generation.

The Department Adviser hands every student a username and password, so accounts are
created in bulk from the class masterlist:

    Student No. | Last Name | First Name | Middle Name | Year Level | Section

* username  = the student number (case-insensitive unique)
* password  = random, readable (no 0/O/1/l/I), shown ONCE on the printable slips

``parse_upload`` -> ``validate_rows`` -> ``import_rows`` are kept free of request/session
handling so they are easy to test.
"""

import csv
import io
import re
import secrets

from django.contrib.auth import get_user_model
from django.contrib.auth.hashers import PBKDF2PasswordHasher
from django.db import transaction

from accounts.models import Student

User = get_user_model()

COLUMNS = ["Student No.", "Last Name", "First Name", "Middle Name", "Year Level", "Section"]
MAX_ROWS = 2000
MAX_FILE_BYTES = 2 * 1024 * 1024

# Header aliases (normalised: lower-case, letters only) -> internal key.
_HEADERS = {
    "studentno": "student_number", "studentnumber": "student_number", "studentid": "student_number",
    "idnumber": "student_number", "idno": "student_number", "id": "student_number",
    "lastname": "last_name", "surname": "last_name", "familyname": "last_name",
    "firstname": "first_name", "givenname": "first_name",
    "middlename": "middle_name", "middleinitial": "middle_name",
    "yearlevel": "year_level", "year": "year_level",
    "section": "section", "block": "section",
}
_YEARS = {
    "1": "1", "1st": "1", "first": "1", "i": "1",
    "2": "2", "2nd": "2", "second": "2", "ii": "2",
    "3": "3", "3rd": "3", "third": "3", "iii": "3",
    "4": "4", "4th": "4", "fourth": "4", "iv": "4",
}
_PASSWORD_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZabcdefghjkmnpqrstuvwxyz23456789"
PASSWORD_LENGTH = 8
# Imports hash hundreds of passwords inside one request, so use a lighter PBKDF2 cost;
# Django re-hashes with the full default cost the first time the student signs in.
_IMPORT_ITERATIONS = 100_000


class MasterlistError(Exception):
    """The uploaded file cannot be read at all (wrong type, empty, no header...)."""


# ---------------------------------------------------------------------------
# Template files
# ---------------------------------------------------------------------------
def template_csv_bytes():
    out = io.StringIO()
    csv.writer(out).writerow(COLUMNS)
    return ("﻿" + out.getvalue()).encode("utf-8")      # BOM so Excel opens it cleanly


def template_xlsx_bytes():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Masterlist"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="41B422")
    for col, width in zip("ABCDEF", (16, 22, 22, 18, 14, 10)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    help_ws = wb.create_sheet("Instructions")
    for line in (
        ["How to fill in the Masterlist sheet"],
        [],
        ["Student No.", "Required, unique. It also becomes the student's username."],
        ["Last Name / First Name", "Required."],
        ["Middle Name", "Optional."],
        ["Year Level", "Optional: 1, 2, 3 or 4 (also accepts 1st / 1st Year ...)."],
        ["Section", "Optional: A, B, C, D, E, F or G."],
        [],
        ["Example", "2024-0001 | Dela Cruz | Juan | Santos | 1st Year | A"],
        ["Delete nothing from the header row. Passwords are generated automatically."],
    ):
        help_ws.append(line)
    help_ws["A1"].font = Font(bold=True, size=13)
    help_ws.column_dimensions["A"].width = 26
    help_ws.column_dimensions["B"].width = 70

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------
def _norm_header(value):
    return re.sub(r"[^a-z]", "", str(value or "").lower())


def _cell(value):
    """Cell -> clean string (Excel gives numbers as floats: 2024.0 -> '2024')."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _rows_from_matrix(matrix):
    """matrix: iterable of lists of raw cells -> list of dicts keyed by internal names."""
    matrix = [list(r) for r in matrix]
    header_idx = next((i for i, r in enumerate(matrix) if any(_cell(c) for c in r)), None)
    if header_idx is None:
        raise MasterlistError("The file is empty.")
    mapping = {}
    for col, raw in enumerate(matrix[header_idx]):
        key = _HEADERS.get(_norm_header(raw))
        if key and key not in mapping.values():
            mapping[col] = key
    missing = [label for label, key in (("Student No.", "student_number"),
                                        ("Last Name", "last_name"),
                                        ("First Name", "first_name"))
               if key not in mapping.values()]
    if missing:
        raise MasterlistError(
            "Missing required column(s): " + ", ".join(missing)
            + ". Download the template and keep its header row."
        )
    rows = []
    for offset, raw in enumerate(matrix[header_idx + 1:], start=header_idx + 2):
        record = {key: _cell(raw[col]) if col < len(raw) else "" for col, key in mapping.items()}
        if not any(record.values()):
            continue                                    # blank line
        record["line"] = offset
        rows.append(record)
    if not rows:
        raise MasterlistError("The file has a header but no students.")
    if len(rows) > MAX_ROWS:
        raise MasterlistError(f"Too many rows ({len(rows)}). Import at most {MAX_ROWS} students at a time.")
    return rows


def parse_upload(uploaded):
    """Read an uploaded .csv / .xlsx file into a list of row dicts."""
    name = (getattr(uploaded, "name", "") or "").lower()
    if getattr(uploaded, "size", 0) > MAX_FILE_BYTES:
        raise MasterlistError("The file is too large (limit 2 MB).")
    data = uploaded.read()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:                       # noqa: BLE001 - any corrupt-file error
            raise MasterlistError("That Excel file could not be read.") from exc
        sheet = wb["Masterlist"] if "Masterlist" in wb.sheetnames else wb.worksheets[0]
        return _rows_from_matrix(sheet.iter_rows(values_only=True))
    if name.endswith(".csv") or name.endswith(".txt"):
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise MasterlistError("That CSV file could not be decoded.")
        sample = text[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return _rows_from_matrix(csv.reader(io.StringIO(text), dialect))
    raise MasterlistError("Upload a .csv or .xlsx file (use the template).")


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------
def normalize_year(value):
    """'1st Year' / 'First' / 1 -> '1'; '' -> ''; unknown -> None."""
    v = re.sub(r"\s*(year|yr)\.?$", "", (value or "").strip().lower()).strip()
    if not v:
        return ""
    return _YEARS.get(v)


#: (key in the row, label, how to read the current value from a Student)
_UPDATABLE = (
    ("last_name", "Last name", lambda s: s.user.last_name),
    ("first_name", "First name", lambda s: s.user.first_name),
    ("middle_name", "Middle name", lambda s: s.user.middle_name),
    ("year_level", "Year", lambda s: s.year_level),
    ("section", "Section", lambda s: s.section),
)


def _changes(student, row):
    """``[(key, label, old, new)]`` the row would change on an existing student.

    A column left blank in the file never erases what is stored (only the two names are
    required, and they are never blank here).
    """
    out = []
    for key, label, current in _UPDATABLE:
        new = (row.get(key) or "").strip()
        old = current(student) or ""
        if new and new != old:
            out.append((key, label, old, new))
    return out


def validate_rows(rows):
    """Annotate each row with ``status`` (new / update / exists / error) and ``message``.

    Rows are normalised in place. The **Student No.** is the key: an unknown number is ``new``;
    a registered one is ``update`` when the file carries a different name / year / section
    (``changes`` lists them) and ``exists`` when nothing differs. Passwords are never touched.
    """
    seen = set()
    existing = {s.student_number.lower(): s for s in Student.objects.select_related("user")}
    existing_usernames = {u.lower() for u in User.objects.values_list("username", flat=True)}
    for row in rows:
        number = row.get("student_number", "")
        problems = []
        if not number:
            problems.append("Student No. is required")
        elif len(number) > 50:
            problems.append("Student No. is longer than 50 characters")
        elif not re.fullmatch(r"[\w.@+\-]+", number):
            problems.append("Student No. may only contain letters, digits and . @ + - _")
        if not row.get("last_name"):
            problems.append("Last Name is required")
        if not row.get("first_name"):
            problems.append("First Name is required")
        year = normalize_year(row.get("year_level", ""))
        if year is None:
            problems.append(f"Year Level '{row.get('year_level')}' is not 1-4")
            year = row.get("year_level", "")     # keep the raw value: confirm re-validates
        row["year_level"] = year
        raw_section = (row.get("section") or "").strip()
        section = Student.normalize_section(raw_section)
        if raw_section and not section:
            problems.append(f"Section '{raw_section}' is not A-G")
            section = raw_section[:20]               # keep the raw value: confirm re-validates
        row["section"] = section
        key = number.lower()
        if problems:
            row["status"], row["message"] = "error", "; ".join(problems)
        elif key in seen:
            row["status"], row["message"] = "error", "Duplicate Student No. in this file"
        elif key in existing:
            changes = _changes(existing[key], row)
            row["changes"] = [f"{label}: {old or '-'} -> {new}" for _, label, old, new in changes]
            if changes:
                row["status"], row["message"] = "update", "; ".join(row["changes"])
            else:
                row["status"], row["message"] = "exists", "Already registered - no changes"
        elif key in existing_usernames:
            row["status"], row["message"] = "error", "That username is already used by another account"
        else:
            row["status"], row["message"] = "new", ""
        if number:
            seen.add(key)
    return rows


def review_order(rows):
    """Rows for the preview: the ones that need attention (errors) first, then file order."""
    return sorted(rows, key=lambda r: (r["status"] != "error", r.get("line", 0)))


def error_report_rows(rows):
    """The error rows as ``COLUMNS`` + ``Line`` + ``Problem`` dicts, ready to fix and upload again."""
    out = []
    for r in rows:
        if r.get("status") != "error":
            continue
        out.append({
            "Student No.": r.get("student_number", ""), "Last Name": r.get("last_name", ""),
            "First Name": r.get("first_name", ""), "Middle Name": r.get("middle_name", ""),
            "Year Level": r.get("year_level", ""), "Section": r.get("section", ""),
            "Line": r.get("line", ""), "Problem": r.get("message", ""),
        })
    return out


def summarize(rows):
    return {s: sum(1 for r in rows if r["status"] == s) for s in ("new", "update", "exists", "error")}


# ---------------------------------------------------------------------------
# Creating accounts / credentials
# ---------------------------------------------------------------------------
def generate_password(length=PASSWORD_LENGTH):
    return "".join(secrets.choice(_PASSWORD_ALPHABET) for _ in range(length))


def _hash(password):
    hasher = PBKDF2PasswordHasher()
    return hasher.encode(password, hasher.salt(), iterations=_IMPORT_ITERATIONS)


def _credential(student, password):
    user = student.user
    return {
        "student_number": student.student_number,
        "name": user.get_full_name() or user.username,
        "year_section": student.year_section,
        "username": user.username,
        "password": password,
    }


def import_rows(rows):
    """Create an account for every ``new`` row and update every ``update`` row.

    Returns the credentials of the created accounts (updated students keep their password).
    """
    credentials = []
    with transaction.atomic():
        for row in rows:
            if row.get("status") == "update":
                student = Student.objects.select_related("user").filter(
                    student_number__iexact=row["student_number"]).first()
                if student is not None:
                    _apply_changes(student, row)
                continue
            if row.get("status") != "new":
                continue
            password = generate_password()
            user = User(
                username=row["student_number"],
                first_name=row["first_name"],
                middle_name=row.get("middle_name", ""),
                last_name=row["last_name"],
                is_student=True,
                is_active=True,
            )
            user.password = _hash(password)
            user.save()
            student = Student.objects.create(
                user=user,
                student_number=row["student_number"],
                year_level=row.get("year_level", ""),
                section=row.get("section", ""),
            )
            credentials.append(_credential(student, password))
    return credentials


def _apply_changes(student, row):
    user_fields, student_fields = [], []
    for key, _label, _old, new in _changes(student, row):
        if key in ("last_name", "first_name", "middle_name"):
            setattr(student.user, key, new)
            user_fields.append(key)
        else:
            setattr(student, key, new)
            student_fields.append(key)
    if user_fields:
        student.user.save(update_fields=user_fields)
    if student_fields:
        student.save(update_fields=student_fields)


def reset_passwords(students):
    """Give each student in the queryset a fresh random password -> credentials list."""
    credentials = []
    with transaction.atomic():
        for student in students.select_related("user"):
            password = generate_password()
            student.user.password = _hash(password)
            student.user.save(update_fields=["password"])
            credentials.append(_credential(student, password))
    return credentials
